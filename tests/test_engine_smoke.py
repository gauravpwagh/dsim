"""Regression smoke test: pins run_simulation()'s output for one corridor against a
golden Excel file captured from the pre-restructure script (final_sim_sj_4aug.py,
run twice and confirmed byte-identical in cell content -- the engine is fully
deterministic given its fixed seeds, see docs/restructure-notes.md).

This is deliberately narrow -- one corridor, cell-value comparison only -- but it is
the first thing that should catch a behavior change from any future refactor of
src/iidsim/engine/simulate.py, especially the deeper decomposition recommended in
restructure-notes.md as a follow-up. Extend it (more corridors, targeted synthetic
scenarios for autoblock/starvation/sibling-redirect/platform-fallback) before
attempting that decomposition.
"""
import pathlib

import openpyxl
import pytest

from iidsim.engine import run_simulation

GOLDEN_DIR = pathlib.Path(__file__).parent / "golden"


def _cells_by_sheet(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return {name: [tuple(r) for r in wb[name].iter_rows(values_only=True)] for name in wb.sheetnames}


def test_sprd_vzm_matches_golden_output(tmp_path):
    result = run_simulation(
        corridor_dataset="p_g_sprd_vzm_2days",
        network_section="sprd_vzm",
        output_dir=str(tmp_path),
    )

    actual = _cells_by_sheet(result["excel_filename"])
    expected = _cells_by_sheet(GOLDEN_DIR / "p_g_sprd_vzm_2days_golden.xlsx")

    assert actual.keys() == expected.keys()
    for sheet_name in expected:
        assert actual[sheet_name] == expected[sheet_name], f"Sheet {sheet_name!r} diverged from golden output"
