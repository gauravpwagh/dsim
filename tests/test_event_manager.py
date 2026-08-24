"""Stage 0 verification for docs/event-manager-design.md: the heap-based EventManager
must reproduce build_event_list()'s behavior exactly, not just "look right." Every
scenario here is run through both code paths (use_event_manager=False, the existing
default, and True, the new path) and the resulting total_schedule -- the complete
record of every simulated arrival/departure -- must match exactly. Any divergence in
event ordering, tie-breaking, or the refresh-candidate set (see run.py's
docstring comment where refresh_candidates is built) would show up here as a
mismatch, not just as a crash.
"""
import pathlib

import openpyxl
import pytest

from iidsim.engine import run_simulation

from .scenarios import (
    build_autoblock_scenario,
    build_platform_fallback_scenario,
    build_starvation_scenario,
    build_tie_break_scenario,
)

GOLDEN_DIR = pathlib.Path(__file__).parent / "golden"


def _run_both(name, trains, tmp_path, network_section="psa_ktv", **kwargs):
    (tmp_path / "old").mkdir()
    (tmp_path / "new").mkdir()
    old = run_simulation(
        corridor_dataset=name, network_section=network_section, trains_override=trains,
        output_dir=str(tmp_path / "old"), use_event_manager=False, **kwargs,
    )
    new = run_simulation(
        corridor_dataset=name, network_section=network_section, trains_override=trains,
        output_dir=str(tmp_path / "new"), use_event_manager=True, **kwargs,
    )
    return old, new


@pytest.mark.parametrize("name, builder, extra_kwargs", [
    ("scn_autoblock", build_autoblock_scenario, {"autoblock_stations": ("smlg", "kvls")}),
    ("scn_starvation", build_starvation_scenario, {"goods_max_priority_wait_hours": 0.02}),
    ("scn_platform", build_platform_fallback_scenario, {}),
    ("scn_tie_break", build_tie_break_scenario, {}),
])
def test_event_manager_matches_build_event_list(name, builder, extra_kwargs, tmp_path):
    old, new = _run_both(
        name, builder(), tmp_path,
        use_halt_deviation=False, use_speed_randomness=False, **extra_kwargs,
    )
    assert new["total_schedule"] == old["total_schedule"], (
        f"{name}: EventManager path diverged from build_event_list() path -- "
        "event ordering, tie-breaking, or the refresh-candidate set is wrong"
    )


def test_event_manager_matches_on_real_corridor_with_randomness(tmp_path):
    """The synthetic scenarios above disable halt/speed randomness for determinism.
    This exercises the EventManager path against a real corridor with randomness ON --
    the actual production configuration -- and separately checks it still matches the
    engine's own golden Excel file, the same one test_engine_smoke.py pins the
    build_event_list() path against."""
    old, new = _run_both("p_g_sprd_vzm_2days", None, tmp_path, network_section="sprd_vzm")
    assert new["total_schedule"] == old["total_schedule"]

    def cells_by_sheet(path):
        wb = openpyxl.load_workbook(path, data_only=True)
        return {s: [tuple(r) for r in wb[s].iter_rows(values_only=True)] for s in wb.sheetnames}

    actual = cells_by_sheet(new["excel_filename"])
    expected = cells_by_sheet(GOLDEN_DIR / "p_g_sprd_vzm_2days_golden.xlsx")
    assert actual.keys() == expected.keys()
    for sheet_name in expected:
        assert actual[sheet_name] == expected[sheet_name], (
            f"EventManager path diverged from golden output on sheet {sheet_name!r}"
        )
