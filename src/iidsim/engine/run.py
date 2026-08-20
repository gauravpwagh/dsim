"""The Simulation class: combines every mixin, the main event loop (run()), resource_update_event() (the core per-event state transition -- by far the largest single method, deliberately kept as one function rather than further split, see docs/restructure-notes.md for why), and the Excel/chart/JSON report-generation helpers.

run_simulation() below is the public entry point: build a Simulation, run it, return the result dict -- unchanged from before this split, just delegating to the class now instead of executing one large function body."""

from iidsim.engine.state import SimulationState
from iidsim.engine.resolve import ResolveMixin
from iidsim.engine.priority import PriorityMixin
from iidsim.engine.randomness import RandomnessMixin
from iidsim.engine.events import EventsMixin
import json
# import time
# import numpy as np
import pandas as pd
# from pandas import Timestamp
# from scipy import stats as _halt_dev_stats
# from openpyxl.styles import Alignment, Font
# import iidsim.network as _network
# from iidsim import schedules
# from iidsim.data import geography, halt_deviation, timing
from iidsim.reporting.chart import plot_railway_chart
from iidsim.reporting.extract import filter_df_by_date_window, get_formatted_data_from_df

class Simulation(SimulationState, ResolveMixin, PriorityMixin, RandomnessMixin, EventsMixin):

    def resource_update_event(self, next_event_type, next_event_tr_id, t):
        stn_line_occ_flag = 0
        stn_line_occ_end = pd.Timestamp('2100-06-01 22:50:00')
        stn_line_occ_endtimes = []
        len_sched = len(self.sched_act[next_event_tr_id])
        t_ind = self.sched_act[next_event_tr_id].index(t)
        if next_event_type == 'a':
            print('\n event at time t = ', t, ' is arrival of train ', self.next_event_train, ' at station', self.stns_event[0].name)
            t_end = self.sched_act[next_event_tr_id][t_ind + 1]
            print('\n scheduled station halt end time is: ', t_end)
            stn_line = self.stn_line_assign(t_ind, next_event_tr_id, len_sched)
            print('\n station line assignment information: ', stn_line)
            stn_line_occ_flag = stn_line[0]
            if stn_line_occ_flag == 0:
                is_origin_stn = t_ind == 1
                if is_origin_stn:
                    min_endtime = self.get_min_endtime(self.stns_event[0], t_ind, t_ind == 1)
                    new_arr_time = min_endtime + pd.Timedelta(minutes=1)
                    if new_arr_time <= t:
                        new_arr_time = t + pd.Timedelta(minutes=1)
                    self.sched_updt(new_arr_time, t_ind, next_event_tr_id)
                    self.t_max = self.term_crit_calc()
                else:
                    min_endtime = self.get_min_endtime(self.stns_event[0], t_ind, t_ind == 1)
                    print('min end time is ', min_endtime)
                    new_arr_time = min_endtime + pd.Timedelta(minutes=1)
                    if new_arr_time <= t:
                        new_arr_time = t + pd.Timedelta(minutes=1)
                    new_occ_inc = new_arr_time - t
                    print('\n new occ increament: ', new_occ_inc)
                    self.sched_updt(new_arr_time, t_ind, next_event_tr_id)
                    self.blsec_t.train_occ_updt(new_arr_time, 1)
                    print('\n block section ', self.blsec_t.name, ' will be occupied by train ', self.next_event_train, 'until ', new_arr_time)
                    if len(self.blsec_t.blsec_queue) > 0:
                        q_len = len(self.blsec_t.blsec_queue)
                        q_tr_id = []
                        v = 0
                        while v < q_len:
                            if self.blsec_t.blsec_queue[v] in self.sched_act:
                                q_tr_id.append(self.blsec_t.blsec_queue[v])
                            v += 6
                        redirected = []
                        _queue_dir = 'up' if self.station_longitudes[self.stns_event[1].name] > self.station_longitudes[self.stns_event[0].name] else 'dn'
                        for k in q_tr_id:
                            train_name_r = k[:k.index('_')]
                            stn_line_stn0_r = self.get_train_stnline_for_arrival_delay(self.stns_event, train_name_r)
                            stn_line_stn1_r = self.get_train_stnline_for_departure_delay(self.stns_event, train_name_r)
                            if stn_line_stn0_r is not None:
                                redirect_stn_obj, redirect_stn_line = (self.stns_event[1], stn_line_stn0_r)
                            elif stn_line_stn1_r is not None:
                                redirect_stn_obj, redirect_stn_line = (self.stns_event[0], stn_line_stn1_r)
                            else:
                                redirect_stn_obj, redirect_stn_line = (None, None)
                            sib = self.find_free_sibling_blsec(self.blsec_t, redirect_stn_obj, redirect_stn_line, train_dir=_queue_dir) if redirect_stn_obj is not None else None
                            if sib is not None:
                                qi = self.blsec_t.blsec_queue.index(k)
                                tr_index_q = self.blsec_t.blsec_queue[qi + 3]
                                print('queued train', k, 'redirected off still-delayed', self.blsec_t.name, '-> free sibling', sib.name, 'is available')
                                self.blsec_t.queue_remove([k])
                                self.sched_updt(t + pd.Timedelta(minutes=1), tr_index_q, k)
                                redirected.append(k)
                        q_tr_id = [k for k in q_tr_id if k not in redirected]
                        self.blsec_t.queue_updt(new_occ_inc)
                        print('\n list of train schedules from queue to be updated: ', q_tr_id)
                        for k in q_tr_id:
                            u = 0
                            for l in self.sched_act[k]:
                                if isinstance(l, str) == False and l < pd.Timestamp('2100-06-01 22:50:00'):
                                    self.sched_act[k][u] = l + new_occ_inc
                                u += 1
                            train_name = k[:k.index('_')]
                            stn_line_stn0 = self.get_train_stnline_for_arrival_delay(self.stns_event, train_name)
                            stn_line_stn1 = self.get_train_stnline_for_departure_delay(self.stns_event, train_name)
                            print('station line is ', stn_line_stn0, stn_line_stn1)
                            if stn_line_stn0 is not None:
                                updt_dep_time = self.stns_event[1].tracks[stn_line_stn0][3] + new_occ_inc
                                self.stns_event[1].set_occupancy_updt(stn_line_stn0, updt_dep_time)
                                print(f'[queue update] train {train_name} at {self.stns_event[1].name}, stn line {stn_line_stn0}, updated occ_end to {updt_dep_time}')
                            elif stn_line_stn1 is not None:
                                updt_dep_time = self.stns_event[0].tracks[stn_line_stn1][3] + new_occ_inc
                                self.stns_event[0].set_occupancy_updt(stn_line_stn1, updt_dep_time)
                                print(f'[queue update] train {train_name} at {self.stns_event[0].name}, stn line {stn_line_stn1}, updated occ_end to {updt_dep_time}')
                            else:
                                print(f'[queue update] train {train_name} not found at either station, skipping station line update')
                        print('\n updated block section queue status: \n', self.blsec_t.blsec_queue)
                    print('\n updated block section queue status after adding current event train at starting: \n', self.blsec_t.blsec_queue)
                    self.t_max = self.term_crit_calc()
                    print('\n updated time at which simulation terminates: ', self.t_max)
                    if not is_origin_stn:
                        if len(self.blsec_t.autoblsec_list) > 0:
                            check_if_auto = self.autoblsec_check()
                            if next_event_tr_id == self.blsec_t.autoblsec_list[0][0] and check_if_auto:
                                self.blsec_t.autoblsecsection_trains_updt(new_occ_inc)
                                print('\n updated autoblock section list for block section ', self.blsec_t.name, ' is now ', self.blsec_t.autoblsec_list)
                                for lst in self.blsec_t.autoblsec_list:
                                    print('\n updating individual train schedule for train ', lst)
                                    tr_id = lst[0]
                                    new_arr_time = lst[4]
                                    print('\n new arrival time after updating autoblock section list is: ', new_arr_time)
                                    new_ind = lst[2]
                                    print('\n index position: ', new_ind, t_ind)
                                    self.sched_updt(pd.Timestamp(new_arr_time), new_ind, tr_id)
                    self.t_max = self.term_crit_calc()
            else:
                self.stn_line_occ_name = stn_line[1]
                original_halt_duration = t_end - t
                Y = self.add_halt_randomness()
                new_halt_duration = original_halt_duration + pd.Timedelta(minutes=Y)
                t_end_new = t + new_halt_duration
                t_end_new = t_end_new.replace(microsecond=0)
                if t_end_new < t:
                    t_end_new = t
                print('\n new station halt duration after adding randomness is: ', new_halt_duration, 'and', Y)
                self.stns_event[0].set_occupancy_arr(self.stn_line_occ_name, t, t_end_new, self.tr_next_event.train_id)
                self.stns_event[0].set_occupancy_updt(stn_line[1], t_end_new)
                self.sched_updt(t_end_new, t_ind + 1, next_event_tr_id)
                self.tr_next_event.tr_sched_act[self.stns_event[0].name][1] = t_end_new
                print('\n train arrival event details at ', self.stns_event[0].name, ' station are:', self.stns_event[0].tracks[stn_line[1]][2])
                print('\n -----> station line occupied: ', self.stn_line_occ_name)
                if self.stns_event[0].tracks[self.stn_line_occ_name][1] != 0:
                    print('\n -----> platform occupied: ', 'P' + str(self.stns_event[0].tracks[self.stn_line_occ_name][1]))
                if t_ind >= 4:
                    print('train arrival blocksection is ', self.blsec_t.name)
                    print('blocksection direction name is ', self.blsec_t.dir_mvmt)
                    next_event_conn = self.blsec_t.name + '_' + self.stn_line_occ_name
                    print('connection name is ', next_event_conn)
                    self.stns_event[0].set_occ_conn_in(next_event_conn, self.tr_next_event.train_id, 0)
                    print('\n station line connection ', next_event_conn, ' is now free', self.stns_event[0].connections[next_event_conn])
                    self.blsec_t.train_occ_end(t)
                    print('\n block section ', self.blsec_t.name, ' is now free')
                    print('\n prev block section (blsec_t) occupancy status is ', self.blsec_t.occ_ind)
                    print('\n prev block section occ end time is ', self.blsec_t.occ_end)
                self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:50:00')
                self.tr_sched_updt(t, next_event_type)
                occupancy_end = self.sched_act[next_event_tr_id][t_ind + 1]
                print('station line ', self.stn_line_occ_name, 'occ end time at station', self.stns_event[0].name, ' is ', self.stns_event[0].tracks[self.stn_line_occ_name][3])
                self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, occupancy_end)
                print('stns evnet length is ', len(self.stns_event))
                if len(self.stns_event) > 1:
                    print('inside to remvoe autoblock train ')
                    print('blocksection name is ', self.blsec_t.name)
                    check_if_auto = self.autoblsec_check()
                    if self.blsec_t.autoblsec_list and next_event_tr_id == self.blsec_t.autoblsec_list[0][0] and check_if_auto:
                        self.blsec_t.autoblsecsection_trains_remove()
                        print('\n train ', self.next_event_train, ' removed from autoblock section list of block section ', self.blsec_t.autoblsec_list)
                self.t_max = self.term_crit_calc()
                print('t_max is ', self.t_max)
                try:
                    platform = self.stns_event[0].tracks[self.stn_line_occ_name][1]
                except Exception:
                    platform = None
                self.total_schedule[next_event_tr_id]['simulated'].append([self.stns_event[0].name, self.stn_line_occ_name, platform])
                self.total_schedule[next_event_tr_id]['simulated'].append([t])
        else:
            print('\n event at time t = ', t, ' is departure of train ', self.next_event_train, ' from station', self.stns_event[0].name)
            self.stn_line_occ_name = ''
            for j in self.stns_event[0].tracks:
                print('\n station line under consideration for departure event: ', j)
                if self.stns_event[0].tracks[j][0] == 1 and self.stns_event[0].tracks[j][5] == self.tr_next_event.train_id:
                    self.stn_line_occ_name = j
                    print('\n station line from which train ', self.next_event_train, ' will depart is ', self.stn_line_occ_name)
                    break
            print('\n current train sched is ', self.sched_act[next_event_tr_id])
            if t_ind == len_sched - 1:
                self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                print('\n station line ', self.stn_line_occ_name, 'at station ', self.stns_event[0].name, 'will be free starting at t = ', t)
                self.tr_sched_updt(t, next_event_type)
                self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                self.total_schedule[next_event_tr_id]['simulated'].append([t])
            if t_ind < len_sched - 1:
                next_event_conn = self.blsec_t.name + '_' + self.stn_line_occ_name
                print('departure event connection is ', next_event_conn)
                occ_details = self.check_next_blse_stn_occupancy(t_ind, len_sched)
                print('occ_details list is ', occ_details)
                stn1 = self.stns_event[0].name
                stn2 = self.stns_event[1].name
                current_train_dir = 1 if self.station_longitudes[stn2] > self.station_longitudes[stn1] else 0
                dn_stn_count = occ_details[0]
                up_stn_count = occ_details[1]
                count_next_blsec = occ_details[3]
                count_curr_blsec = occ_details[4]
                next_blsec_list = occ_details[5]
                total_stn_lines = len(self.stns_event[1].tracks)
                total_stn_occ = dn_stn_count + up_stn_count
                same_dir_stn_count = dn_stn_count if current_train_dir == 0 else up_stn_count
                free_stn_lines = total_stn_lines - total_stn_occ
                curr_is_single = count_curr_blsec == 1
                next_is_single = count_next_blsec == 1
                ready_to_dept = True
                if curr_is_single and next_is_single:
                    print('both current and next block sections are single line')
                    if free_stn_lines >= 2:
                        ready_to_dept = True
                    elif free_stn_lines == 1:
                        print('one stn line is free block')
                        blsec_obj, blsec_dir = next_blsec_list[0]
                        print('next block section occupancy status is ', next_blsec_list[0])
                        print('dir and block is', blsec_dir, current_train_dir, blsec_obj.occ_ind)
                        if blsec_obj.occ_ind == 0:
                            ready_to_dept = True
                        elif blsec_dir == current_train_dir and blsec_obj.occ_ind == 1:
                            ready_to_dept = True
                        else:
                            ready_to_dept = False
                    elif free_stn_lines == 0:
                        if same_dir_stn_count == 0:
                            ready_to_dept = False
                        else:
                            blsec_obj, blsec_dir = next_blsec_list[0]
                            if blsec_obj.occ_ind == 0:
                                ready_to_dept = True
                            elif blsec_dir == current_train_dir:
                                ready_to_dept = True
                            else:
                                ready_to_dept = False
                elif curr_is_single and (not next_is_single):
                    print('current block section is single line and next block section is double line')
                    if free_stn_lines >= 2:
                        ready_to_dept = True
                    elif free_stn_lines == 1:
                        free_blsec_count = sum((1 for b in next_blsec_list if b[0].occ_ind == 0))
                        same_dir_blsec_count = sum((1 for b in next_blsec_list if b[1] == current_train_dir))
                        if free_blsec_count == len(next_blsec_list):
                            ready_to_dept = True
                        elif free_blsec_count >= 1 and same_dir_stn_count >= 1:
                            ready_to_dept = True
                        elif same_dir_blsec_count >= 1:
                            ready_to_dept = True
                        else:
                            print('the only free station line is occupied by an oncoming train and both lines of the next block section are occupied by oncoming trains, so we need to wait at the station')
                            ready_to_dept = False
                    elif free_stn_lines == 0 and same_dir_stn_count == 0:
                        print('all station lines are occupied by opposing trains, so we need to wait at the station')
                        ready_to_dept = False
                    elif free_stn_lines == 0 and same_dir_stn_count >= 1:
                        same_dir_blsec_count = sum((1 for b in next_blsec_list if b[1] == current_train_dir))
                        if same_dir_blsec_count >= 1:
                            ready_to_dept = True
                        else:
                            print('all station lines are occupied by opposing trains, but at least one train has same direction as given train and both lines of the next block section are occupied by oncoming trains, so we need to wait at the station')
                            ready_to_dept = False
                elif not curr_is_single and next_is_single:
                    print('current block section is double line and next block section is single line')
                    if free_stn_lines >= 2:
                        ready_to_dept = True
                    elif free_stn_lines == 1:
                        blsec_obj, blsec_dir = next_blsec_list[0]
                        if blsec_obj.occ_ind == 0:
                            ready_to_dept = True
                        elif blsec_dir == current_train_dir:
                            ready_to_dept = True
                        else:
                            print('the only free station line is occupied by an oncoming train and the next block section is single line occupied by an oncoming train, so we need to wait at the station')
                            ready_to_dept = False
                    elif free_stn_lines == 0:
                        if same_dir_stn_count == 0:
                            print('all station lines are occupied by opposing trains, so we need to wait at the station')
                            ready_to_dept = False
                        else:
                            blsec_obj, blsec_dir = next_blsec_list[0]
                            if blsec_obj.occ_ind == 0:
                                ready_to_dept = True
                            elif blsec_dir == current_train_dir:
                                ready_to_dept = True
                            else:
                                print('the only free station line is occupied by an oncoming train and the next block section is single line occupied by an oncoming train, so we need to wait at the station')
                                ready_to_dept = False
                elif ready_to_dept is False:
                    self.t_max = self.term_crit_calc()
                else:
                    ready_to_dept = True
                    print('single line block section case is not found ready to depart True')
                print(f'ready_to_dept={ready_to_dept} | free_stn={free_stn_lines} | same_dir_stn={same_dir_stn_count} | curr_SL={curr_is_single} | next_SL={next_is_single}')
                check_if_auto = self.autoblsec_check()
                print('autoblock case is result is ', check_if_auto)
                if self.blsec_t.occ_ind == 0 and len(self.blsec_t.blsec_queue) == 0 and (not check_if_auto) and (ready_to_dept is True):
                    print('\n case when block free and queue is empty')
                    if t_ind >= 4 and self.tr_next_event.tr_type == 'g':
                        updt_dep_time = self.goods_delay_due_to_passenger(self.sched_act, t, t_ind, next_event_tr_id)
                        original_dep_time = self.sched_act[next_event_tr_id][t_ind]
                        starved_too_long = t - self.given_sched[next_event_tr_id][t_ind] >= self.GOODS_MAX_PRIORITY_WAIT
                        if starved_too_long:
                            print('goods train', next_event_tr_id, 'has waited', t - self.given_sched[next_event_tr_id][t_ind], '-> overriding passenger priority, departing now')
                        if updt_dep_time != original_dep_time and (not self.check_if_all_goods(t_ind, next_event_tr_id)) and (not starved_too_long):
                            self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, updt_dep_time + pd.Timedelta(minutes=1))
                            self.sched_updt(updt_dep_time + pd.Timedelta(minutes=1), t_ind, next_event_tr_id)
                            self.t_max = self.term_crit_calc()
                            print('After consideration of passenger trains at current station, and that coming from prev blsec /n')
                        else:
                            new_arrival_time = self.new_arr_by_speed_randomness(next_event_tr_id, t_ind)[0]
                            self.sched_updt(new_arrival_time, t_ind + 2, next_event_tr_id)
                            t_blsec_occ_end = new_arrival_time - t
                            self.t_max = self.term_crit_calc()
                            _expected_dir = 'up' if self.station_longitudes[self.stns_event[0].name] > self.station_longitudes[self.stns_event[1].name] else 'dn'
                            assert self.blsec_t.dir_mvmt[0:2] == _expected_dir or self.blsec_t.dir_mvmt[0:3] == 'mid', f'direction mismatch: train {next_event_tr_id} departing {self.stns_event[0].name}->{self.stns_event[1].name} (expected dir {_expected_dir!r}) about to occupy {self.blsec_t.name} (dir_mvmt={self.blsec_t.dir_mvmt!r})'
                            self.blsec_t.train_occ_start(t, t + t_blsec_occ_end, next_event_tr_id)
                            print('\n block section ', self.blsec_t.name, ' will be occupied by train ', self.next_event_train, ' until ', new_arrival_time)
                            self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                            print('\n train ', self.next_event_train, 'departing station ', self.stns_event[0].name, ' from connection ', next_event_conn)
                            self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                            print('\n station line ', self.stn_line_occ_name, 'at station ', self.stns_event[0].name, 'will be free starting at t = ', t)
                            self.tr_sched_updt(t, next_event_type)
                            self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                            print('\n block section ', self.blsec_t.name, 'queue status: ', self.blsec_t.blsec_queue)
                            self.total_schedule[next_event_tr_id]['simulated'].append([t])
                    else:
                        new_arrival_time = self.new_arr_by_speed_randomness(next_event_tr_id, t_ind)[0]
                        self.sched_updt(new_arrival_time, t_ind + 2, next_event_tr_id)
                        t_blsec_occ_end = new_arrival_time - t
                        self.t_max = self.term_crit_calc()
                        _expected_dir = 'up' if self.station_longitudes[self.stns_event[0].name] > self.station_longitudes[self.stns_event[1].name] else 'dn'
                        assert self.blsec_t.dir_mvmt[0:2] == _expected_dir or self.blsec_t.dir_mvmt[0:3] == 'mid', f'direction mismatch: train {next_event_tr_id} departing {self.stns_event[0].name}->{self.stns_event[1].name} (expected dir {_expected_dir!r}) about to occupy {self.blsec_t.name} (dir_mvmt={self.blsec_t.dir_mvmt!r})'
                        self.blsec_t.train_occ_start(t, t + t_blsec_occ_end, next_event_tr_id)
                        print('\n block section ', self.blsec_t.name, ' will be occupied by train ', self.next_event_train, ' until ', new_arrival_time)
                        self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                        print('\n train ', self.next_event_train, 'departing station ', self.stns_event[0].name, ' from connection ', next_event_conn)
                        self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                        print('\n station line ', self.stn_line_occ_name, 'at station ', self.stns_event[0].name, 'will be free starting at t = ', t)
                        self.tr_sched_updt(t, next_event_type)
                        self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                        print('\n block section ', self.blsec_t.name, 'queue status: ', self.blsec_t.blsec_queue)
                        self.total_schedule[next_event_tr_id]['simulated'].append([t])
                elif self.blsec_t.occ_ind == 0 and len(self.blsec_t.blsec_queue) > 0 and (not check_if_auto) and (ready_to_dept is True):
                    print('when block is free and queue is not empty')
                    if self.tr_next_event.tr_type == 'g' and t_ind >= 4:
                        updt_dep_time = self.goods_delay_due_to_passenger(self.sched_act, t, t_ind, next_event_tr_id)
                        original_dep_time = self.sched_act[next_event_tr_id][t_ind]
                        starved_too_long = t - self.given_sched[next_event_tr_id][t_ind] >= self.GOODS_MAX_PRIORITY_WAIT
                        if starved_too_long:
                            print('goods train', next_event_tr_id, 'has waited', t - self.given_sched[next_event_tr_id][t_ind], '-> overriding passenger priority, departing now')
                        if updt_dep_time != original_dep_time and (not self.check_if_all_goods(t_ind, next_event_tr_id)) and (not starved_too_long):
                            self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, updt_dep_time + pd.Timedelta(minutes=1))
                            self.sched_updt(updt_dep_time + pd.Timedelta(minutes=1), t_ind, next_event_tr_id)
                            self.t_max = self.term_crit_calc()
                            print('After consideration of passenger trains at current station, and that coming from prev blsec /n')
                            time_increament = updt_dep_time + pd.Timedelta(minutes=1) - t
                            self.blsec_t.queue_updt(time_increament)
                            for i in range(0, len(self.blsec_t.blsec_queue), 6):
                                trainid, traintype, blsec_occupancy_start = (self.blsec_t.blsec_queue[i], self.blsec_t.blsec_queue[i + 2], self.blsec_t.blsec_queue[i + 4])
                                for line, data in self.stns_event[0].tracks.items():
                                    print('given train id is ', trainid)
                                    print('goods train data is ', data)
                                    print('goods train station line', line)
                                    if data[-1] == trainid:
                                        self.stns_event[0].set_occupancy_updt(line, blsec_occupancy_start + pd.Timedelta(minutes=1))
                                        self.sched_updt(blsec_occupancy_start + pd.Timedelta(minutes=1), self.blsec_t.blsec_queue[i + 3], trainid)
                                        self.t_max = self.term_crit_calc()
                        else:
                            speed_rand = self.new_arr_by_speed_randomness(next_event_tr_id, t_ind)
                            new_arrival_time = speed_rand[0]
                            if len(speed_rand) > 1 and len(self.blsec_t.blsec_queue) > 6:
                                time_inc = speed_rand[1]
                                q_len = len(self.blsec_t.blsec_queue)
                                q_tr_id = []
                                v = 0
                                while v < q_len:
                                    if self.blsec_t.blsec_queue[v] in self.sched_act and self.blsec_t.blsec_queue[v] != next_event_tr_id:
                                        q_tr_id.append([self.blsec_t.blsec_queue[v], self.blsec_t.blsec_queue[v + 3]])
                                    v += 6
                                self.blsec_t.queue_updt(time_inc)
                                print('list of train schedules from queue to be updated due to speed randomness: ', q_tr_id)
                                for k, t_ind_q in q_tr_id:
                                    self.sched_updt(self.sched_act[k][t_ind_q] + time_inc, t_ind_q, k)
                                    train_name = k[:k.index('_')]
                                    stn_line_stn0 = self.get_train_stnline_for_arrival_delay(self.stns_event, train_name)
                                    stn_line_stn1 = self.get_train_stnline_for_departure_delay(self.stns_event, train_name)
                                    if stn_line_stn0 is not None:
                                        updt_dep_time = self.stns_event[1].tracks[stn_line_stn0][3] + time_inc
                                        self.stns_event[1].set_occupancy_updt(stn_line_stn0, updt_dep_time)
                                        print(f'[queue update] train {train_name} at {self.stns_event[1].name}, stn line {stn_line_stn0}, updated occ_end to {updt_dep_time}')
                                    elif stn_line_stn1 is not None:
                                        updt_dep_time = self.stns_event[0].tracks[stn_line_stn1][3] + time_inc
                                        self.stns_event[0].set_occupancy_updt(stn_line_stn1, updt_dep_time)
                                        print(f'[queue update] train {train_name} at {self.stns_event[0].name}, stn line {stn_line_stn1}, updated occ_end to {updt_dep_time}')
                                    else:
                                        print(f'[queue update] train {train_name} not found at either station, skipping station line update')
                                print('updated block section queue status due to speed randomness: ', self.blsec_t.blsec_queue)
                            self.sched_updt(new_arrival_time, t_ind + 2, next_event_tr_id)
                            t_blsec_occ_end = new_arrival_time - t
                            self.t_max = self.term_crit_calc()
                            _expected_dir = 'up' if self.station_longitudes[self.stns_event[0].name] > self.station_longitudes[self.stns_event[1].name] else 'dn'
                            assert self.blsec_t.dir_mvmt[0:2] == _expected_dir or self.blsec_t.dir_mvmt[0:3] == 'mid', f'direction mismatch: train {next_event_tr_id} departing {self.stns_event[0].name}->{self.stns_event[1].name} (expected dir {_expected_dir!r}) about to occupy {self.blsec_t.name} (dir_mvmt={self.blsec_t.dir_mvmt!r})'
                            _expected_dir = 'up' if self.station_longitudes[self.stns_event[0].name] > self.station_longitudes[self.stns_event[1].name] else 'dn'
                            assert self.blsec_t.dir_mvmt[0:2] == _expected_dir or self.blsec_t.dir_mvmt[0:3] == 'mid', f'direction mismatch: train {next_event_tr_id} departing {self.stns_event[0].name}->{self.stns_event[1].name} (expected dir {_expected_dir!r}) about to occupy {self.blsec_t.name} (dir_mvmt={self.blsec_t.dir_mvmt!r})'
                            self.blsec_t.train_occ_start(t, t + t_blsec_occ_end, next_event_tr_id)
                            print('\n block section ', self.blsec_t.name, ' will be occupied by train ', self.next_event_train, ' until ', new_arrival_time)
                            self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                            print('\n train ', self.next_event_train, 'departing station ', self.stns_event[0].name, ' from connection ', next_event_conn)
                            self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                            print('\n station line ', self.stn_line_occ_name, 'at station ', self.stns_event[0].name, 'will be free starting at t = ', t)
                            self.tr_sched_updt(t, next_event_type)
                            self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                            t_index = t_ind
                            dep_list = [next_event_tr_id, self.next_event_train, self.tr_next_event.tr_type, t_index, t, self.sched_act[next_event_tr_id][t_ind + 2]]
                            self.blsec_t.queue_remove(dep_list)
                            print('\n block section ', self.blsec_t.name, 'queue status: ', self.blsec_t.blsec_queue)
                            print('current train sched is ', self.sched_act[next_event_tr_id])
                            self.total_schedule[next_event_tr_id]['simulated'].append([t])
                    else:
                        speed_rand = self.new_arr_by_speed_randomness(next_event_tr_id, t_ind)
                        new_arrival_time = speed_rand[0]
                        if len(speed_rand) > 1 and len(self.blsec_t.blsec_queue) > 6:
                            time_inc = speed_rand[1]
                            q_len = len(self.blsec_t.blsec_queue)
                            q_tr_id = []
                            v = 0
                            while v < q_len:
                                if self.blsec_t.blsec_queue[v] in self.sched_act and self.blsec_t.blsec_queue[v] != next_event_tr_id:
                                    q_tr_id.append([self.blsec_t.blsec_queue[v], self.blsec_t.blsec_queue[v + 3]])
                                v += 6
                            self.blsec_t.queue_updt(time_inc)
                            print('list of train schedules from queue to be updated due to speed randomness: ', q_tr_id)
                            for k, t_ind_q in q_tr_id:
                                self.sched_updt(self.sched_act[k][t_ind_q] + time_inc, t_ind_q, k)
                                train_name = k[:k.index('_')]
                                stn_line_stn0 = self.get_train_stnline_for_arrival_delay(self.stns_event, train_name)
                                stn_line_stn1 = self.get_train_stnline_for_departure_delay(self.stns_event, train_name)
                                if stn_line_stn0 is not None:
                                    updt_dep_time = self.stns_event[1].tracks[stn_line_stn0][3] + time_inc
                                    self.stns_event[1].set_occupancy_updt(stn_line_stn0, updt_dep_time)
                                    print(f'[queue update] train {train_name} at {self.stns_event[1].name}, stn line {stn_line_stn0}, updated occ_end to {updt_dep_time}')
                                elif stn_line_stn1 is not None:
                                    updt_dep_time = self.stns_event[0].tracks[stn_line_stn1][3] + time_inc
                                    self.stns_event[0].set_occupancy_updt(stn_line_stn1, updt_dep_time)
                                    print(f'[queue update] train {train_name} at {self.stns_event[0].name}, stn line {stn_line_stn1}, updated occ_end to {updt_dep_time}')
                                else:
                                    print(f'[queue update] train {train_name} not found at either station, skipping station line update')
                            print('updated block section queue status due to speed randomness: ', self.blsec_t.blsec_queue)
                        self.sched_updt(new_arrival_time, t_ind + 2, next_event_tr_id)
                        print('new arrival time is ', new_arrival_time)
                        print('after schedule update for arrival time, schedule is ', self.sched_act[next_event_tr_id])
                        t_blsec_occ_end = new_arrival_time - t
                        self.t_max = self.term_crit_calc()
                        _expected_dir = 'up' if self.station_longitudes[self.stns_event[0].name] > self.station_longitudes[self.stns_event[1].name] else 'dn'
                        assert self.blsec_t.dir_mvmt[0:2] == _expected_dir or self.blsec_t.dir_mvmt[0:3] == 'mid', f'direction mismatch: train {next_event_tr_id} departing {self.stns_event[0].name}->{self.stns_event[1].name} (expected dir {_expected_dir!r}) about to occupy {self.blsec_t.name} (dir_mvmt={self.blsec_t.dir_mvmt!r})'
                        self.blsec_t.train_occ_start(t, t + t_blsec_occ_end, next_event_tr_id)
                        print('\n block section ', self.blsec_t.name, ' will be occupied by train ', self.next_event_train, ' until ', t + t_blsec_occ_end)
                        self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                        print('\n train ', self.next_event_train, 'departing station ', self.stns_event[0].name, ' from connection ', next_event_conn)
                        print(self.stns_event[0].tracks[self.stn_line_occ_name][0] if self.stn_line_occ_name in self.stns_event[0].tracks else 'N/A (train has no station track)')
                        self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                        print('\n station line ', self.stn_line_occ_name, 'at station ', self.stns_event[0].name, 'will be free starting at t = ', t)
                        print(self.stns_event[0].tracks[self.stn_line_occ_name][0] if self.stn_line_occ_name in self.stns_event[0].tracks else 'N/A (train has no station track)')
                        self.tr_sched_updt(t, next_event_type)
                        self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                        t_index = t_ind
                        dep_list = [next_event_tr_id, self.next_event_train, self.tr_next_event.tr_type, t_index, t, self.sched_act[next_event_tr_id][t_ind + 2]]
                        self.blsec_t.queue_remove(dep_list)
                        print('\n block section ', self.blsec_t.name, 'queue status: ', self.blsec_t.blsec_queue)
                        self.total_schedule[next_event_tr_id]['simulated'].append([t])
                elif check_if_auto:
                    print('auto block section name is ', self.blsec_t.name)
                    if (self.stns_event[0].name and self.stns_event[1].name) in self.autoblock_stations:
                        print('inside the autoblock section case')
                        if len(self.blsec_t.autoblsec_list) < 1:
                            start_time = t
                            end_time = self.sched_act[next_event_tr_id][t_ind + 2]
                            print('blsec_t.length=', self.blsec_t.length)
                            print('duration=', end_time - start_time)
                            speed = self.blsec_t.length / ((end_time - start_time).total_seconds() / 3600)
                            print('calculated speed for autoblock section is ', speed)
                            self.blsec_t.autoblsecsection_trains([next_event_tr_id, speed, t_ind + 2, start_time, end_time])
                            self.tr_sched_updt(t, next_event_type)
                            self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                            occupancy_end = self.sched_act[next_event_tr_id][t_ind + 1]
                            self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                            self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                            self.total_schedule[next_event_tr_id]['simulated'].append([t])
                            print('autoblock section case list', self.blsec_t.autoblsec_list)
                        elif len(self.blsec_t.autoblsec_list) >= 1:
                            last_train_speed = self.blsec_t.autoblsec_list[-1][1]
                            print('blsec_t.length=', self.blsec_t.length)
                            print('duration=', self.sched_act[next_event_tr_id][t_ind + 2] - t)
                            current_train_speed = self.blsec_t.length / ((self.sched_act[next_event_tr_id][t_ind + 2] - t).total_seconds() / 3600)
                            print('last train speed in autoblock section is ', last_train_speed, 'current train speed is ', current_train_speed)
                            time_taken_by_last_train = 3.6 / last_train_speed
                            print('time taken by last train to cover 3.6km safe distance is ', time_taken_by_last_train, ' hr')
                            print(pd.Timedelta(hours=time_taken_by_last_train), pd.Timedelta(hours=time_taken_by_last_train).total_seconds() / 60)
                            print(self.blsec_t.autoblsec_list[-1][3], type(self.blsec_t.autoblsec_list[-1][3]))
                            last_train_time_to_safe_distance = self.blsec_t.autoblsec_list[-1][3] + pd.Timedelta(hours=time_taken_by_last_train)
                            print('delay in minutes is', pd.Timedelta(hours=time_taken_by_last_train).total_seconds() / 60)
                            print('last train time to safe distance is ', last_train_time_to_safe_distance)
                            if self.tr_next_event.tr_type == 'p':
                                if last_train_speed > current_train_speed:
                                    print('inside the condition when last train speed > current train speed')
                                    if t >= last_train_time_to_safe_distance:
                                        blsec_start_time = t
                                        end_time = self.sched_act[next_event_tr_id][t_ind + 2]
                                        speed = self.blsec_t.length / ((end_time - t).total_seconds() / 3600)
                                        self.blsec_t.autoblsecsection_trains([next_event_tr_id, speed, t_ind + 2, blsec_start_time, end_time])
                                        self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                                        occupancy_end = self.sched_act[next_event_tr_id][t_ind + 1]
                                        self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                                        self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                                        self.tr_sched_updt(t, next_event_type)
                                        print('autoblock section list', self.blsec_t.autoblsec_list)
                                        self.total_schedule[next_event_tr_id]['simulated'].append([t])
                                    else:
                                        print('Case of departure delay')
                                        dept_time = last_train_time_to_safe_distance
                                        self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, dept_time)
                                        self.sched_updt(dept_time, t_ind, next_event_tr_id)
                                        self.tr_sched_updt(dept_time, next_event_type)
                                        print('autoblock section list', self.blsec_t.autoblsec_list)
                                elif last_train_speed == current_train_speed:
                                    print('inside the condition when last train speed = current train speed')
                                    if t >= last_train_time_to_safe_distance:
                                        blsec_start_time = t
                                        end_time = self.sched_act[next_event_tr_id][t_ind + 2]
                                        speed = self.blsec_t.length / ((end_time - t).total_seconds() / 3600)
                                        self.blsec_t.autoblsecsection_trains([next_event_tr_id, speed, t_ind + 2, blsec_start_time, end_time])
                                        self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                                        occupancy_end = self.sched_act[next_event_tr_id][t_ind + 1]
                                        self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                                        self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                                        self.tr_sched_updt(t, next_event_type)
                                        print('autoblock section list', self.blsec_t.autoblsec_list)
                                        self.total_schedule[next_event_tr_id]['simulated'].append([t])
                                    else:
                                        dept_time = last_train_time_to_safe_distance
                                        self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, dept_time)
                                        self.sched_updt(dept_time, t_ind, next_event_tr_id)
                                        self.tr_sched_updt(t, next_event_type)
                                        print('autoblock section list', self.blsec_t.autoblsec_list)
                                elif last_train_speed < current_train_speed:
                                    print('inside the condition when last train spped < current train speed')
                                    t_d1 = pd.Timestamp(self.blsec_t.autoblsec_list[-1][3])
                                    print('t_d1 value is ', t_d1)
                                    print('t_d1 is which is previous trains dept time is', t_d1)
                                    print('type of t_d1 is', type(t_d1))
                                    t_d2 = t_d1 + pd.Timedelta(hours=self.blsec_t.length / last_train_speed - (self.blsec_t.length - self.headway_distance) / current_train_speed)
                                    print('calculated t_d2 is ', t_d2)
                                    print('type of t_d2 is', type(t_d2))
                                    speed = self.blsec_t.length / ((self.sched_act[next_event_tr_id][t_ind + 2] - t).total_seconds() / 3600)
                                    print('calculated speed for autoblock section is ', speed)
                                    if t_d2 <= t:
                                        print('case when t_d2 <= t')
                                        blsec_start_time = t
                                        self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:50:00')
                                        end_time = self.sched_act[next_event_tr_id][t_ind + 2]
                                        self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                                        self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                                        self.blsec_t.autoblsecsection_trains([next_event_tr_id, speed, t_ind + 2, blsec_start_time, end_time])
                                        self.tr_sched_updt(t, next_event_type)
                                        print('train added in autoblock section list is ', self.blsec_t.autoblsec_list)
                                        self.total_schedule[next_event_tr_id]['simulated'].append([t])
                                    else:
                                        print('case when t_d2 > t; means train got delayed')
                                        dept_time = t_d2
                                        self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, dept_time)
                                        self.sched_updt(dept_time, t_ind, next_event_tr_id)
                                        self.tr_sched_updt(dept_time, next_event_type)
                            if self.tr_next_event.tr_type == 'g':
                                print('Given train is Goods')
                                if last_train_speed > current_train_speed:
                                    print('inside the condition when last train speed > current train speed')
                                    if t >= last_train_time_to_safe_distance:
                                        blsec_start_time = t
                                        end_time = self.sched_act[next_event_tr_id][t_ind + 2]
                                        speed = self.blsec_t.length / ((end_time - t).total_seconds() / 3600)
                                        self.blsec_t.autoblsecsection_trains([next_event_tr_id, speed, t_ind + 2, blsec_start_time, end_time])
                                        self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                                        occupancy_end = self.sched_act[next_event_tr_id][t_ind + 1]
                                        self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                                        self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                                        self.tr_sched_updt(t, next_event_type)
                                        print('autoblock section list', self.blsec_t.autoblsec_list)
                                        self.total_schedule[next_event_tr_id]['simulated'].append([t])
                                    else:
                                        dept_time = last_train_time_to_safe_distance
                                        self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, dept_time)
                                        self.sched_updt(dept_time, t_ind, next_event_tr_id)
                                        self.tr_sched_updt(dept_time, next_event_type)
                                        print('autoblock section list', self.blsec_t.autoblsec_list)
                                elif last_train_speed == current_train_speed:
                                    print('inside the condition when last train speed = current train speed')
                                    if t >= last_train_time_to_safe_distance:
                                        blsec_start_time = t
                                        end_time = self.sched_act[next_event_tr_id][t_ind + 2]
                                        speed = self.blsec_t.length / ((end_time - t).total_seconds() / 3600)
                                        self.blsec_t.autoblsecsection_trains([next_event_tr_id, speed, t_ind + 2, blsec_start_time, end_time])
                                        self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:52:00')
                                        occupancy_end = self.sched_act[next_event_tr_id][t_ind + 1]
                                        self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                                        self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                                        self.tr_sched_updt(t, next_event_type)
                                        print('autoblock section list', self.blsec_t.autoblsec_list)
                                        self.total_schedule[next_event_tr_id]['simulated'].append([t])
                                    else:
                                        dept_time = last_train_time_to_safe_distance
                                        self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, dept_time)
                                        self.sched_updt(dept_time, t_ind, next_event_tr_id)
                                        self.tr_sched_updt(t, next_event_type)
                                        print('autoblock section list', self.blsec_t.autoblsec_list)
                                elif last_train_speed < current_train_speed:
                                    print('inside the condition when last train spped < current train speed')
                                    t_d1 = pd.Timestamp(self.blsec_t.autoblsec_list[-1][3])
                                    print('t_d1 is which is previous trains dept time is', t_d1)
                                    print('type of t_d1 is', type(t_d1))
                                    t_d2 = t_d1 + pd.Timedelta(hours=self.blsec_t.length / last_train_speed - (self.blsec_t.length - self.headway_distance) / current_train_speed)
                                    print('calculated t_d2 is ', t_d2)
                                    print('type of t_d2 is', type(t_d2))
                                    speed = self.blsec_t.length / ((self.sched_act[next_event_tr_id][t_ind + 2] - t).total_seconds() / 3600)
                                    print('calculated speed for autoblock section is ', speed)
                                    if t_d2 <= t:
                                        print('case when t_d2 <= t')
                                        blsec_start_time = t
                                        self.sched_act[next_event_tr_id][t_ind] = pd.Timestamp('2100-06-01 22:50:00')
                                        end_time = self.sched_act[next_event_tr_id][t_ind + 2]
                                        self.stns_event[0].set_occupancy_dep(self.stn_line_occ_name, t)
                                        self.stns_event[0].set_occ_conn_out(next_event_conn, self.tr_next_event.train_id, 1)
                                        self.blsec_t.autoblsecsection_trains([next_event_tr_id, speed, t_ind + 2, blsec_start_time, end_time])
                                        self.tr_sched_updt(t, next_event_type)
                                        print('train added in autoblock section list is ', self.blsec_t.autoblsec_list)
                                        self.total_schedule[next_event_tr_id]['simulated'].append([t])
                                    else:
                                        print('case when t_d2 > t; means train got delayed')
                                        dept_time = t_d2
                                        self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, dept_time)
                                        self.sched_updt(dept_time, t_ind, next_event_tr_id)
                                        self.tr_sched_updt(dept_time, next_event_type)
                elif ready_to_dept is False:
                    new_halt_end = occ_details[2] + pd.Timedelta(minutes=2)
                    if new_halt_end <= t:
                        new_halt_end = t + pd.Timedelta(minutes=1)
                    self.sched_updt(new_halt_end, t_ind, next_event_tr_id)
                    self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, new_halt_end)
                    if next_event_tr_id in self.blsec_t.blsec_queue[0::6]:
                        q_idx = self.blsec_t.blsec_queue.index(next_event_tr_id)
                        self.blsec_t.blsec_queue[q_idx + 4] = new_halt_end
                        self.blsec_t.blsec_queue[q_idx + 5] = self.sched_act[next_event_tr_id][t_ind + 2]
                        print('refreshed stale queue entry for', next_event_tr_id, 'in', self.blsec_t.name)
                    print('after updating train shced ', self.sched_act[next_event_tr_id])
                else:
                    print('\n Third case when blocksection is not free \n ')
                    print('before adding or updating queue; queue status is ', self.blsec_t.blsec_queue)
                    if self.blsec_t.occ_ind == 1 and len(self.blsec_t.blsec_queue) < 6:
                        t_dep_updt = self.blsec_t.occ_end + pd.Timedelta(minutes=1)
                        print(t_dep_updt)
                    elif next_event_tr_id not in self.blsec_t.blsec_queue:
                        last_queued_occ_end = self.blsec_t.blsec_queue[-1]
                        t_dep_updt = max(last_queued_occ_end + pd.Timedelta(minutes=1), self.blsec_t.occ_end + pd.Timedelta(minutes=1))
                    else:
                        t_dep_updt = self.blsec_t.occ_end + pd.Timedelta(minutes=1)
                    if t_dep_updt <= t:
                        t_dep_updt = t + pd.Timedelta(minutes=1)
                    print('\n block section ', self.blsec_t.name, ' is not free at t = ', t)
                    print('\n block section currently occupied by ', self.blsec_t.occ_train)
                    print('\n current block section is ', self.blsec_t.name)
                    base = '_'.join(self.blsec_t.name.split('_')[:-1])
                    for suffix in ('up1', 'dn1'):
                        sib_name = f'{base}_{suffix}'
                        sib = self.blsec_lookup.get(sib_name)
                        if sib is not None:
                            print(sib.name, '| occ_train:', sib.occ_train, '| occ_end:', sib.occ_end)
                    print('\n block section ', self.blsec_t.name, ' will become free for ', next_event_tr_id, ' at t = ', t_dep_updt - pd.Timedelta(minutes=1))
                    self.stns_event[0].set_occupancy_updt(self.stn_line_occ_name, t_dep_updt)
                    self.sched_updt(t_dep_updt, t_ind, next_event_tr_id)
                    print('current train sched is ', self.sched_act[next_event_tr_id])
                    blsec_occ_end = self.sched_act[next_event_tr_id][t_ind + 2]
                    if next_event_tr_id not in self.blsec_t.blsec_queue:
                        t_index = t_ind
                        self.blsec_t.queue_add(next_event_tr_id, self.next_event_train, self.tr_next_event.tr_type, t_index, t_dep_updt, blsec_occ_end)
                        print('After queue_add:', self.blsec_t.name, '  ', self.blsec_t.blsec_queue)
                        self.update_blsec_queue_priority(self.blsec_t, self.tr_next_event.tr_type, next_event_tr_id)
                        i = 0
                        while i < len(self.blsec_t.blsec_queue):
                            stn_line = self.get_train_stnline_for_departure_delay(self.stns_event, self.blsec_t.blsec_queue[i + 1])
                            if stn_line is not None:
                                self.stns_event[0].set_occupancy_updt(stn_line, self.blsec_t.blsec_queue[i + 4])
                            self.sched_updt(self.blsec_t.blsec_queue[i + 4], self.blsec_t.blsec_queue[i + 3], self.blsec_t.blsec_queue[i])
                            i += 6
                        print('\n after update_blsec_queue_priority:', self.blsec_t.name, '  ', self.blsec_t.blsec_queue)
                        print('\n given blsec_t occ_end is ', self.blsec_t.occ_end)
                    else:
                        print('time diff is', t_dep_updt - t)
                        self.blsec_t.queue_updt(t_dep_updt - t)
                        q_len = len(self.blsec_t.blsec_queue)
                        q_tr_id = []
                        v = 0
                        while v < q_len:
                            if self.blsec_t.blsec_queue[v] in self.sched_act:
                                q_tr_id.append([self.blsec_t.blsec_queue[v], self.blsec_t.blsec_queue[v + 3]])
                            v += 6
                        print('\n list of train schedules from queue to be updated: ', q_tr_id)
                        for k in q_tr_id:
                            train_name = k[0][:k[0].index('_')]
                            tr_index = k[1]
                            print('index number is ', tr_index)
                            print('next evetn train id is ', next_event_tr_id)
                            print('current train id is ', k[0])
                            stn_line = self.get_train_stnline_for_departure_delay(self.stns_event, train_name)
                            print('train no is ', train_name, 'and its station line is ', stn_line)
                            print('train id is ', k[0])
                            train_id = k[0]
                            print('train schedule before updation is ', self.sched_act[train_id])
                            if stn_line is not None:
                                updt_dep_time = self.stns_event[0].tracks[stn_line][3] + (t_dep_updt - t)
                                print('updt dept time for a given train is ', updt_dep_time)
                                self.stns_event[0].set_occupancy_updt(stn_line, updt_dep_time)
                                self.sched_updt(updt_dep_time, tr_index, k[0])
                            else:
                                print('train', train_name, 'not at any station track, skipping station occupancy update')
                    print('t_dep_updt is', self.sched_act[next_event_tr_id][t_ind])
                    print('\n station line occupied by the given train is ', self.stns_event[0].tracks.get(self.stn_line_occ_name, 'N/A (train has no station track)'))
                    self.t_max = self.term_crit_calc()
                    print('\n updated time at which simulation terminates: ', self.t_max)
                    print('\n block section ', self.blsec_t.name, 'queue status: ', self.blsec_t.blsec_queue)

    def run(self):
        while self.t <= self.t_max and self.exec_sim == 0:
            event_list = self.build_event_list()
            next_event_time = min([i for i in event_list if isinstance(i, str) == False and i < pd.Timestamp('2100-06-01 22:50:00')])
            _event_idx = event_list.index(next_event_time)
            self.next_event_type = event_list[_event_idx + 1]
            self.next_event_train = event_list[_event_idx - 1]
            self.next_event_tr_id = event_list[_event_idx + 2]
            self.t = next_event_time
            print('\n -------------------t = ', self.t, '-------------------------')
            print('\n event list at t = ', self.t, ' ', event_list)
            print('\n next event type is ', self.next_event_type)
            self.tr_next_event = self.trains_by_instance_id[self.next_event_tr_id]
            self.stns_event = self.get_station_event(self.t, self.next_event_tr_id, self.next_event_type, self.stations_list)
            if self.next_event_type == 'd':
                if len(self.stns_event) > 1:
                    self.blsec_t = self.blsec_id(self.stns_event[0].name, self.stns_event[1].name)
                else:
                    self.blsec_t = None
            elif len(self.stns_event) > 1:
                self.blsec_t = self.blsec_id(self.stns_event[1].name, self.stns_event[0].name)
            else:
                self.blsec_t = None
            if self.blsec_t is not None:
                print('\n current block section is ', self.blsec_t.name)
                base = '_'.join(self.blsec_t.name.split('_')[:-1])
                for suffix in ('up1', 'dn1'):
                    sib_name = f'{base}_{suffix}'
                    sib = self.blsec_lookup.get(sib_name)
                    if sib is not None:
                        print(sib.name, '| occ_train:', sib.occ_train, '| occ_end:', sib.occ_end)
            self.resource_update_event(self.next_event_type, self.next_event_tr_id, self.t)
            self.t_max = self.term_crit_calc()
        print('whole train schedule is ', self.next_event_tr_id, self.tr_next_event.tr_sched_act)
        print('whole sched is really ', self.sched_act[self.next_event_tr_id])
        print('\n ----------------simulation has ended-------------------------')
        for idx, j in enumerate(self.trains, start=1):
            print(f'Train serial number: {idx}')
            print('\n planned schedule for train', j.train_id, 'is: \n', j.tr_schedule)
            print('\n schedule operated for train', j.train_id, 'is: \n', j.tr_sched_act)
            print('\n train performance statistics: \n', j.calc_tr_stats())
            print('\n ---------------------next train--------------------------')
        rows = []
        max_len = 0
        for train_id, data in self.total_schedule.items():
            for sched_type, sched_data in data.items():
                formatted = [', '.join(map(str, item)) for item in sched_data]
                train_type = formatted[0]
                row = [train_id, sched_type, train_type] + formatted[1:]
                rows.append(row)
                max_len = max(max_len, len(row))
        for r in rows:
            r += [''] * (max_len - len(r))
        self.columns = ['Train_ID', 'Schedule_Type', 'Train_Type']
        n_extra = max_len - 3
        pattern = ['Stn', 'Arr', 'Dept']
        for i in range(n_extra):
            name = pattern[i % 3] + str(i // 3 + 1)
            self.columns.append(name)
        df = pd.DataFrame(rows, columns=self.columns)
        max_allowed_deviation = pd.Timedelta(minutes=5)
        extra_rows = []
        self.time_cols = [c for c in self.columns if c.startswith('Arr') or c.startswith('Dept')]
        for train_id in df['Train_ID'].unique():
            train_rows = df[df['Train_ID'] == train_id]
            planned_row = train_rows[train_rows['Schedule_Type'] == 'planned']
            simulated_row = train_rows[train_rows['Schedule_Type'] == 'simulated']
            actual_row = train_rows[train_rows['Schedule_Type'] == 'actual']
            if planned_row.empty or simulated_row.empty:
                continue
            if not actual_row.empty:
                extra_rows.extend(self.build_deviation_rows(train_id, simulated_row, actual_row, 'sim vs actual'))
            extra_rows.extend(self.build_deviation_rows(train_id, planned_row, simulated_row, 'sim vs planned'))
        df_extra = pd.DataFrame(extra_rows, columns=self.columns)
        df = pd.concat([df, df_extra], ignore_index=True)
        df['Arr1_sort'] = pd.to_datetime(df['Arr1'], errors='coerce')
        df_sorted = pd.concat([df[df['Train_ID'] == tid] for tid in df.groupby('Train_ID')['Arr1_sort'].first().sort_values().index])
        df_sorted = df_sorted.drop(columns=['Arr1_sort']).reset_index(drop=True)
        planned_completion = max((tr.tr_schedule[tr.tr_destination][0] for tr in self.trains))
        simulated_completion = max((tr.tr_sched_act[tr.tr_destination][0] for tr in self.trains))
        actual_completion_candidates = [tr.tr_real_schedule[tr.tr_destination][0] for tr in self.trains if tr.tr_destination in tr.tr_real_schedule and tr.tr_real_schedule[tr.tr_destination][0] is not None]
        actual_completion = max(actual_completion_candidates) if actual_completion_candidates else None
        dev_summary = []
        for train_id in df['Train_ID'].unique():
            train_rows = df[df['Train_ID'] == train_id]
            planned_row = train_rows[train_rows['Schedule_Type'] == 'planned']
            simulated_row = train_rows[train_rows['Schedule_Type'] == 'simulated']
            actual_row = train_rows[train_rows['Schedule_Type'] == 'actual']
            if planned_row.empty or simulated_row.empty:
                continue
            train_type = planned_row.iloc[0]['Train_Type']
            max_dev_sim_vs_planned = self.get_max_dev(planned_row, simulated_row, self.time_cols)
            max_dev_actual_vs_planned = self.get_max_dev(planned_row, actual_row, self.time_cols) if not actual_row.empty else None
            dev_summary.append({'train_id': train_id, 'train_type': train_type, 'max_dev_actual_vs_planned': max_dev_actual_vs_planned, 'max_dev_sim_vs_planned': max_dev_sim_vs_planned})
        dev_summary_df = pd.DataFrame(dev_summary, columns=['train_id', 'train_type', 'max_dev_actual_vs_planned', 'max_dev_sim_vs_planned'])
        p_actual_str = self.avg_sd_str(dev_summary_df.loc[dev_summary_df.train_type == 'p', 'max_dev_actual_vs_planned'])
        p_actual_str_no_outliers = self.avg_sd_str_no_outliers(dev_summary_df.loc[dev_summary_df.train_type == 'p', 'max_dev_actual_vs_planned'])
        g_actual_str = self.avg_sd_str(dev_summary_df.loc[dev_summary_df.train_type == 'g', 'max_dev_actual_vs_planned'])
        g_actual_str_no_outliers = self.avg_sd_str_no_outliers(dev_summary_df.loc[dev_summary_df.train_type == 'g', 'max_dev_actual_vs_planned'])
        p_sim_str = self.avg_sd_str(dev_summary_df.loc[dev_summary_df.train_type == 'p', 'max_dev_sim_vs_planned'])
        g_sim_str = self.avg_sd_str(dev_summary_df.loc[dev_summary_df.train_type == 'g', 'max_dev_sim_vs_planned'])
        self.stn_cols = [c for c in self.columns if c.startswith('Stn')]
        self.arr_cols = [c for c in self.columns if c.startswith('Arr')]
        self.dept_cols = [c for c in self.columns if c.startswith('Dept')]
        speed_summary = []
        for train_id in df['Train_ID'].unique():
            train_rows = df[df['Train_ID'] == train_id]
            planned_row = train_rows[train_rows['Schedule_Type'] == 'planned']
            simulated_row = train_rows[train_rows['Schedule_Type'] == 'simulated']
            actual_row = train_rows[train_rows['Schedule_Type'] == 'actual']
            if planned_row.empty:
                continue
            train_type = planned_row.iloc[0]['Train_Type']
            if train_type != 'g':
                continue
            speed_summary.append({'train_id': train_id, 'speed_planned': self.compute_train_avg_speed(planned_row), 'speed_actual': self.compute_train_avg_speed(actual_row) if not actual_row.empty else None, 'speed_simulated': self.compute_train_avg_speed(simulated_row) if not simulated_row.empty else None})
        speed_summary_df = pd.DataFrame(speed_summary, columns=['train_id', 'speed_planned', 'speed_actual', 'speed_simulated'])
        g_speed_planned_str = self.avg_sd_str(speed_summary_df['speed_planned'])
        g_speed_actual_str = self.avg_sd_str(speed_summary_df['speed_actual'])
        g_speed_simulated_str = self.avg_sd_str(speed_summary_df['speed_simulated'])
        summary_rows = [['Metric', 'planned', 'actual', 'simulated'], ['overall schedule completion time', str(planned_completion), str(actual_completion) if actual_completion is not None else '', str(simulated_completion)], ['passenger trains: avg absolute deviation from overall schedule (SD, minutes)', '--', p_actual_str, p_sim_str], ['goods trains: avg absolute deviation from overall schedule (SD, minutes)', '--', g_actual_str, g_sim_str], ['average speed of goods trains for the schedule (SD, km/hr)', g_speed_planned_str, g_speed_actual_str, g_speed_simulated_str], ['', '', '', ''], ['', '', '', ''], ['after outliers removal from actual scheduled data', '', '', ''], ['passenger trains: avg absolute deviation from overall schedule (SD, minutes)', '--', p_actual_str_no_outliers, p_sim_str], ['goods trains: avg absolute deviation from overall schedule (SD, minutes)', '--', g_actual_str_no_outliers, g_sim_str]]
        summary_df = pd.DataFrame(summary_rows[1:], columns=summary_rows[0])
        p_sim = self.dev_list(dev_summary_df, 'p', 'max_dev_sim_vs_planned')
        g_sim = self.dev_list(dev_summary_df, 'g', 'max_dev_sim_vs_planned')
        p_act_w = self.dev_list(dev_summary_df, 'p', 'max_dev_actual_vs_planned')
        g_act_w = self.dev_list(dev_summary_df, 'g', 'max_dev_actual_vs_planned')
        p_act_no = self.remove_outliers_iqr(dev_summary_df.loc[dev_summary_df.train_type == 'p', 'max_dev_actual_vs_planned']).tolist()
        g_act_no = self.remove_outliers_iqr(dev_summary_df.loc[dev_summary_df.train_type == 'g', 'max_dev_actual_vs_planned']).tolist()
        series = [p_sim, g_sim, p_act_w, g_act_w, p_act_no, g_act_no]
        max_n = max((len(s) for s in series), default=0)
        title_row = ['', 'Deviations from the planned schedule', '', '', '', '', '']
        row_cmp = ['', 'planned vs simulated', 'planned vs simulated', 'Planned vs actual', 'Planned vs actual', 'Planned vs actual', 'Planned vs actual']
        row_out = ['', '', '', 'with outliers', 'with outliers', 'without outliers', 'without outliers']
        row_type = ['', 'Passenger', 'Goods', 'Passenger', 'Goods', 'Passenger', 'Goods']
        range_row = ['range='] + [f'[{min(s):.2f}, {max(s):.2f}]' if s else '' for s in series]
        value_rows = [[''] + [self.r2(series[c][r]) if r < len(series[c]) else None for c in range(6)] for r in range(max_n)]
        sheet3_grid = [title_row, row_cmp, row_out, row_type, range_row] + value_rows
        from openpyxl.styles import Alignment, Font
        with pd.ExcelWriter(self.excel_filename, engine='openpyxl') as writer:
            df_sorted.to_excel(writer, sheet_name='Sheet1', index=False)
            summary_df.to_excel(writer, sheet_name='Sheet2', index=False)
            ws3 = writer.book.create_sheet('Sheet3')
            for row in sheet3_grid:
                ws3.append(['' if v is None else v for v in row])
            for rng in ('B1:G1', 'B2:C2', 'D2:G2', 'B3:C3', 'D3:E3', 'F3:G3'):
                ws3.merge_cells(rng)
            for r in (1, 2, 3, 4, 5):
                for c in range(1, 8):
                    cell = ws3.cell(row=r, column=c)
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.font = Font(bold=True)
            for r in range(6, 5 + max_n + 1):
                for c in range(2, 8):
                    cell = ws3.cell(row=r, column=c)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '0.00'
        print(f'Saved -> {self.excel_filename} (Sheet1 + Sheet2 + Sheet3)')
        sim_rows = []
        for train_id, data in self.total_schedule.items():
            sched_data = data['simulated']
            formatted = [', '.join(map(str, item)) for item in sched_data]
            train_type = formatted[0]
            row = [train_id, train_type] + formatted[1:]
            sim_rows.append(row)
        max_len = max((len(r) for r in sim_rows))
        for r in sim_rows:
            r += [''] * (max_len - len(r))
        columns_sim = ['Train_ID', 'Train_Type']
        n_extra = max_len - 2
        pattern = ['Stn', 'Arr', 'Dept']
        for i in range(n_extra):
            columns_sim.append(pattern[i % 3] + str(i // 3 + 1))
        df_simulated = pd.DataFrame(sim_rows, columns=columns_sim)
        self.stn_cols = [c for c in df_simulated.columns if c.startswith('Stn')]
        for col in self.stn_cols:
            df_simulated[col] = df_simulated[col].apply(lambda x: x.split(',')[0].strip() if isinstance(x, str) and x != '' else x)
        print(f'Total trains: {len(df_simulated)}')
        self.master_time_distance_chart(df_simulated, self.station_order, self.segments, start_dt=self.start_dt, end_dt=self.start_dt + pd.Timedelta(hours=self.chart_duration_hrs), filename=self.chart_filename)
        import json
        self.total_schedule_to_json(self.total_schedule, self.animator)
        return {'excel_filename': self.excel_filename, 'chart_filename': self.chart_filename, 'animator_filename': self.animator, 'total_schedule': self.total_schedule, 'trains': self.trains}

    def avg_sd_str(self, series):
        s = series.dropna()
        if s.empty:
            return ''
        return f'{s.mean():.1f} ({s.std():.1f})'

    def avg_sd_str_no_outliers(self, series):
        s = self.remove_outliers_iqr(series)
        if s.empty:
            return ''
        return f'{s.mean():.1f} ({s.std():.1f})'

    def build_deviation_rows(self, train_id, row_a, row_b, label):
        """Compare row_a vs row_b across time_cols, return [abs_dev_row, minmax_row]."""
        abs_dev_row = [train_id, f'Abs Deviation: {label}', '']
        abs_devs = []
        for col in self.columns[3:]:
            if col in self.time_cols:
                a_val = row_a.iloc[0][col]
                b_val = row_b.iloc[0][col]
                if pd.isna(a_val) or pd.isna(b_val) or a_val == '' or (b_val == ''):
                    abs_dev_row.append('')
                    continue
                try:
                    a_ts = pd.Timestamp(a_val)
                    b_ts = pd.Timestamp(b_val)
                    dev = abs((b_ts - a_ts).total_seconds() / 60)
                    abs_dev_row.append(f'{dev:.1f} min')
                    abs_devs.append(dev)
                except Exception:
                    abs_dev_row.append('')
            else:
                abs_dev_row.append('')
        if abs_devs:
            min_dev = min(abs_devs)
            max_dev = max(abs_devs)
            minmax_label = f'[{min_dev:.1f}, {max_dev:.1f}] min'
        else:
            minmax_label = ''
        minmax_row = [train_id, f'[Abs Min dev, Abs Max dev]: {label}', minmax_label] + [''] * (len(self.columns) - 3)
        return [abs_dev_row, minmax_row]

    def clean_stn(self, val):
        if val is None or (isinstance(val, float) and pd.isna(val)) or val == '':
            return None
        return str(val).split(',')[0].strip()

    def compute_train_avg_speed(self, row):
        total_dist = 0.0
        total_hours = 0.0
        n = len(self.stn_cols)
        for i in range(n - 1):
            stn_a = self.clean_stn(row.iloc[0][self.stn_cols[i]])
            stn_b = self.clean_stn(row.iloc[0][self.stn_cols[i + 1]])
            dept_a, arr_b = (row.iloc[0][self.dept_cols[i]], row.iloc[0][self.arr_cols[i + 1]])
            if not stn_a or not stn_b or pd.isna(dept_a) or pd.isna(arr_b) or (dept_a == '') or (arr_b == ''):
                continue
            dist = self.get_distance(stn_a, stn_b)
            if dist is None:
                continue
            try:
                hours = (pd.Timestamp(arr_b) - pd.Timestamp(dept_a)).total_seconds() / 3600
            except Exception:
                continue
            if hours <= 0:
                continue
            total_dist += dist
            total_hours += hours
        if total_hours == 0:
            return None
        return total_dist / total_hours

    def dev_list(self, df_, type_, col):
        return df_.loc[df_.train_type == type_, col].dropna().tolist()

    def get_distance(self, stn_a, stn_b):
        if (stn_a, stn_b) in self.block_section_distances:
            return self.block_section_distances[stn_a, stn_b]
        if (stn_b, stn_a) in self.block_section_distances:
            return self.block_section_distances[stn_b, stn_a]
        return None

    def get_max_dev(self, row_a, row_b, time_cols):
        """Max absolute deviation (minutes) across stations between two schedule rows for one train."""
        devs = []
        for col in time_cols:
            a_val = row_a.iloc[0][col]
            b_val = row_b.iloc[0][col]
            if pd.isna(a_val) or pd.isna(b_val) or a_val == '' or (b_val == ''):
                continue
            try:
                a_ts = pd.Timestamp(a_val)
                b_ts = pd.Timestamp(b_val)
                devs.append(abs((b_ts - a_ts).total_seconds() / 60))
            except Exception:
                continue
        return max(devs) if devs else None

    def master_time_distance_chart(self, df, station_order, segments, start_dt, end_dt, filename):
        df_day = filter_df_by_date_window(df, start_dt)
        if len(df_day) == 0:
            print('Stopping -- no data for provided date')
        else:
            train_data, chart_date = get_formatted_data_from_df(df_day, start_dt=start_dt, end_dt=end_dt)
            plot_railway_chart(station_order, segments, train_data, chart_date, filename)

    def r2(self, v):
        return round(float(v), 2) if v is not None else None

    def remove_outliers_iqr(self, series):
        """Drop values outside [Q1 - 1.5*IQR, Q3 + 1.5*IQR]."""
        s = series.dropna()
        if s.empty:
            return s
        q1, q3 = (s.quantile(0.25), s.quantile(0.75))
        iqr = q3 - q1
        lower, upper = (q1 - 1.5 * iqr, q3 + 1.5 * iqr)
        return s[(s >= lower) & (s <= upper)]

    def total_schedule_to_json(self, total_schedule, output_filename='schedule_simulated.json'):
        trains_list = []
        for train_id, data in total_schedule.items():
            sched_data = data['simulated']
            train_type = sched_data[0][0]
            stops = []
            i = 1
            while i + 2 <= len(sched_data) - 1:
                stn_info = sched_data[i]
                station = stn_info[0]
                line = stn_info[1] if len(stn_info) > 1 else 's1'
                platform = stn_info[2] if len(stn_info) > 2 else 1
                next_block = stn_info[3] if len(stn_info) > 3 else None
                arr_str = sched_data[i + 1][0].strftime('%Y-%m-%d %H:%M:%S')
                dep_str = sched_data[i + 2][0].strftime('%Y-%m-%d %H:%M:%S')
                stops.append({'station': station, 'line': line, 'platform': platform, 'arr': arr_str, 'dep': dep_str, 'next_block': next_block})
                i += 3
            route = []
            for idx, stop in enumerate(stops):
                stop_entry = {'station': stop['station'], 'line': stop['line'], 'platform': stop['platform'], 'arr': stop['arr'], 'dep': stop['dep']}
                if idx + 1 < len(stops):
                    outgoing_block = stops[idx + 1]['next_block']
                    if outgoing_block is not None:
                        stop_entry['next_block'] = outgoing_block
                route.append(stop_entry)
            trains_list.append({'train_id': train_id, 'train_type': train_type, 'route': route})
        all_times = []
        for t in trains_list:
            for r in t['route']:
                all_times.append(r['arr'])
                all_times.append(r['dep'])
        output = {'simulation': {'start_time': min(all_times), 'end_time': max(all_times)}, 'trains': trains_list}
        with open(output_filename, 'w') as f:
            json.dump(output, f, indent=2)
        print(f'Saved {len(trains_list)} trains to {output_filename}')

def run_simulation(
    corridor_dataset,
    network_section,
    *,
    trains_override=None,
    output_dir="output_files",
    start_dt_mode="planned",
    start_dt_manual=None,
    start_dt_buffer_minutes=10,
    chart_duration_hrs=23,
    headway_distance=3.6,
    goods_max_priority_wait_hours=20,
    autoblock_stations=("alm", "kuk", "vzm"),
    use_halt_deviation=True,
    use_speed_randomness=True,
    halt_deviation_seed=1234,
    speed_randomness_seed=1234,
):
    """Run one end-to-end simulation and write the Excel report, time-distance
    chart PDF, and animator JSON for `corridor_dataset` on `network_section`.

    corridor_dataset: a name from iidsim.schedules.available_corridors(), e.g.
        'p_g_sprd_vzm_2days' -- replaces the old hardcoded per-corridor import. Still
        used to derive output filenames even when `trains_override` is given.
    network_section: one of 'psa_ktv', 'sprd_vzm', 'krdl_ktv' -- which corridor's
        station order / segment distances to use for the output chart.
    trains_override: optional list[iidsim.domain.train] to simulate directly instead of
        loading `corridor_dataset` from iidsim.schedules -- for synthetic/targeted test
        scenarios built against the real network (see tests/scenarios.py).

    Returns a dict with the three output file paths plus the in-memory
    total_schedule and trains the run produced.
    """
    sim = Simulation(
        corridor_dataset,
        network_section,
        trains_override=trains_override,
        output_dir=output_dir,
        start_dt_mode=start_dt_mode,
        start_dt_manual=start_dt_manual,
        start_dt_buffer_minutes=start_dt_buffer_minutes,
        chart_duration_hrs=chart_duration_hrs,
        headway_distance=headway_distance,
        goods_max_priority_wait_hours=goods_max_priority_wait_hours,
        autoblock_stations=autoblock_stations,
        use_halt_deviation=use_halt_deviation,
        use_speed_randomness=use_speed_randomness,
        halt_deviation_seed=halt_deviation_seed,
        speed_randomness_seed=speed_randomness_seed,
    )
    return sim.run()
